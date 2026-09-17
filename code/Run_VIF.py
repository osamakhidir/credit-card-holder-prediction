import pandas as pd
import numpy as np
from statsmodels.stats.outliers_influence import variance_inflation_factor
from statsmodels.tools.tools import add_constant
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import warnings
warnings.filterwarnings("ignore")

# ── 1. Load & prepare features ─────────────────────────────────────────────
df = pd.read_csv("banking_dataset_sanitized.csv")

drop_cols = [
    "customer_id",              # identifier
    "credit_card_holder_flag",  # target
    "account_open_date",        # string date
    "zip_code",                 # nominal integer
]
numeric_cols = df.select_dtypes(include=["int64", "float64"]).columns.tolist()
feature_cols = [c for c in numeric_cols if c not in drop_cols]

X = df[feature_cols].dropna()
X_const = add_constant(X)

# ── 2. Compute VIF ─────────────────────────────────────────────────────────
vif_data = pd.DataFrame({
    "feature": X_const.columns,
    "VIF"    : [variance_inflation_factor(X_const.values, i)
                for i in range(X_const.shape[1])]
})
vif_data = (
    vif_data[vif_data["feature"] != "const"]
    .assign(VIF=lambda d: d["VIF"].round(2))
    .sort_values("VIF", ascending=False)
    .reset_index(drop=True)
)

def vif_flag(v):
    if v == float("inf") or v >= 10: return "HIGH"
    if v >= 5:                        return "MODERATE"
    return                                   "OK"

vif_data["severity"] = vif_data["VIF"].apply(vif_flag)
vif_data["rank"]     = range(1, len(vif_data) + 1)
vif_data = vif_data[["rank", "feature", "VIF", "severity"]]

# ── 3. Print to console ────────────────────────────────────────────────────
pd.set_option("display.max_rows", None)
print(vif_data.to_string(index=False))
print(f"\n  HIGH     (VIF ≥ 10) : {(vif_data['severity']=='HIGH').sum()}")
print(f"  MODERATE (VIF 5–10) : {(vif_data['severity']=='MODERATE').sum()}")
print(f"  OK       (VIF < 5)  : {(vif_data['severity']=='OK').sum()}")

# ── 4. Write to Excel ──────────────────────────────────────────────────────
RED_FILL    = PatternFill("solid", fgColor="FFCCCC")
AMBER_FILL  = PatternFill("solid", fgColor="FFF2CC")
GREEN_FILL  = PatternFill("solid", fgColor="D9EAD3")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
TITLE_FILL  = PatternFill("solid", fgColor="2E75B6")
GREY_FILL   = PatternFill("solid", fgColor="F2F2F2")

thin       = Side(style="thin", color="BFBFBF")
border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)

wb  = openpyxl.Workbook()
ws1 = wb.active
ws1.title = "VIF Results"

# Title
ws1.merge_cells("A1:D1")
ws1["A1"].value     = "Variance Inflation Factor (VIF) Analysis — banking_dataset_sanitized.csv"
ws1["A1"].fill      = TITLE_FILL
ws1["A1"].font      = Font(name="Arial", bold=True, color="FFFFFF", size=13)
ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[1].height = 28

# Sub-title
ws1.merge_cells("A2:D2")
ws1["A2"].value     = (f"Features assessed: {len(vif_data)}   |   "
                       "Target excluded: credit_card_holder_flag   |   "
                       "Also excluded: customer_id, account_open_date, zip_code")
ws1["A2"].fill      = GREY_FILL
ws1["A2"].font      = Font(name="Arial", italic=True, size=9, color="595959")
ws1["A2"].alignment = Alignment(horizontal="center")
ws1.row_dimensions[2].height = 16

# Headers
for col_idx, h in enumerate(["Rank", "Feature", "VIF Score", "Severity"], 1):
    cell            = ws1.cell(row=3, column=col_idx, value=h)
    cell.fill       = HEADER_FILL
    cell.font       = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    cell.alignment  = Alignment(horizontal="center", vertical="center")
    cell.border     = border_thin
ws1.row_dimensions[3].height = 20

# Data rows
for row_idx, row in vif_data.iterrows():
    excel_row   = row_idx + 4
    sev         = row["severity"]
    row_fill    = RED_FILL if sev == "HIGH" else (AMBER_FILL if sev == "MODERATE" else GREEN_FILL)
    vif_display = "∞" if row["VIF"] == float("inf") else row["VIF"]

    for col_idx, (val, aln) in enumerate(
        zip([row["rank"], row["feature"], vif_display, sev],
            ["center", "left", "center", "center"]), 1
    ):
        cell           = ws1.cell(row=excel_row, column=col_idx, value=val)
        cell.fill      = row_fill
        cell.font      = Font(name="Arial", size=10)
        cell.alignment = Alignment(horizontal=aln, vertical="center")
        cell.border    = border_thin
    ws1.row_dimensions[excel_row].height = 16

ws1.column_dimensions["A"].width = 8
ws1.column_dimensions["B"].width = 32
ws1.column_dimensions["C"].width = 14
ws1.column_dimensions["D"].width = 16
ws1.freeze_panes = "A4"

wb.save("VIF_Analysis.xlsx")
print("\n✅ Saved: VIF_Analysis.xlsx")