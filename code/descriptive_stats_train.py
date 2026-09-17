"""
Descriptive Statistics — train_set.csv
Outputs: descriptive_stats_train_output.xlsx  (6 styled sheets)

Usage (run from folder containing train_set.csv):
    python descriptive_stats_train.py
"""

import os
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import warnings
warnings.filterwarnings("ignore")

# ── Paths ─────────────────────────────────────────────────────────────────────
BASE_DIR    = os.path.dirname(os.path.abspath(__file__))
_candidates = [os.path.join(BASE_DIR, "train_set.csv"), os.path.join(os.getcwd(), "train_set.csv")]
INPUT_FILE  = next((p for p in _candidates if os.path.exists(p)), _candidates[0])
OUTPUT_FILE = os.path.join(os.path.dirname(INPUT_FILE), "descriptive_stats_train_output.xlsx")

# ── Colour palette ────────────────────────────────────────────────────────────
NAVY       = "1F3864"
STEEL      = "2E75B6"
LIGHT_BLUE = "D6E4F0"
WHITE      = "FFFFFF"
DARK_GRAY  = "404040"
GREEN_FILL = "D5F5E3"
RED_FILL   = "FADBD8"
YELLOW_FILL= "FEF9E7"

# ── Style helpers ─────────────────────────────────────────────────────────────
def solid(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def hdr_font(size=10, color=WHITE, bold=True):
    return Font(name="Arial", bold=bold, color=color, size=size)

def body_font(bold=False, color=DARK_GRAY, size=9):
    return Font(name="Arial", bold=bold, color=color, size=size)

def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

def center_align(wrap=True):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def right_align():
    return Alignment(horizontal="right", vertical="center")

def left_align():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def write_banner(ws, row, ncols, title, fill=NAVY, size=12):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=title)
    c.fill = solid(fill); c.font = hdr_font(size=size)
    c.alignment = center_align(); c.border = thin_border()
    ws.row_dimensions[row].height = 30

def write_col_headers(ws, row, headers, fill=STEEL):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = solid(fill); c.font = hdr_font(size=9)
        c.alignment = center_align(); c.border = thin_border()
    ws.row_dimensions[row].height = 36

def w_num(ws, row, col, value, fmt="#,##0.00", fill_hex=None):
    c = ws.cell(row=row, column=col, value=value)
    c.number_format = fmt
    c.font = body_font(); c.alignment = right_align(); c.border = thin_border()
    if fill_hex:
        c.fill = solid(fill_hex)
    return c

def w_str(ws, row, col, value, bold=False, fill_hex=None, align="left"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = body_font(bold=bold); c.border = thin_border()
    c.alignment = center_align() if align == "center" else left_align()
    if fill_hex:
        c.fill = solid(fill_hex)
    return c

def row_fill(r):
    return LIGHT_BLUE if r % 2 == 0 else WHITE

# ── Load data ─────────────────────────────────────────────────────────────────
print("Loading train_set.csv …")
df = pd.read_csv(INPUT_FILE)
print(f"  {df.shape[0]:,} rows × {df.shape[1]} columns")

TARGET = "credit_card_holder_flag"

# Column groups
# Binary: only 0/1 values (includes one-hot dummies + flags)
binary_cols = [c for c in df.columns
               if set(df[c].dropna().unique()).issubset({0, 1, 0.0, 1.0})]

# One-hot dummy prefixes
DUMMY_PREFIXES = [
    "gender_", "education_level_", "employment_status_", "marital_status_",
    "primary_account_type_", "account_status_", "loyalty_tier_",
    "risk_tier_", "state_"
]
dummy_cols    = [c for c in binary_cols
                 if any(c.startswith(p) for p in DUMMY_PREFIXES)]
flag_cols     = [c for c in binary_cols if c not in dummy_cols and c != TARGET]
continuous_cols = [c for c in df.columns if c not in binary_cols]

print(f"  Continuous : {len(continuous_cols)}  |  Flags : {len(flag_cols)}  "
      f"|  Dummies : {len(dummy_cols)}  |  Target : {TARGET}")

# ── Workbook ──────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
wb.remove(wb.active)

# =============================================================================
# SHEET 1 — OVERVIEW
# =============================================================================
ws = wb.create_sheet("Overview")
ws.sheet_view.showGridLines = False

write_banner(ws, 1, 2, "Train Set — Descriptive Statistics Overview")

meta = [
    ("Source file",          "train_set.csv"),
    ("Total rows",           f"{df.shape[0]:,}"),
    ("Total columns",        df.shape[1]),
    ("Continuous columns",   len(continuous_cols)),
    ("Binary flag columns",  len(flag_cols)),
    ("One-hot dummy columns",len(dummy_cols)),
    ("Target column",        TARGET),
    ("Target = 1  (card holders)", int(df[TARGET].sum())),
    ("Target = 0  (non-holders)",  int((df[TARGET] == 0).sum())),
    ("Target imbalance ratio",     f"{(df[TARGET] == 0).sum() / df[TARGET].sum():.1f} : 1"),
    ("Missing values (total)",     int(df.isnull().sum().sum())),
    ("Duplicate rows",             int(df.duplicated().sum())),
]
for i, (label, value) in enumerate(meta, 3):
    fill = LIGHT_BLUE if i % 2 == 0 else WHITE
    c1 = ws.cell(row=i, column=1, value=label)
    c1.font = body_font(bold=True); c1.border = thin_border()
    c1.fill = solid(fill); c1.alignment = left_align()
    c2 = ws.cell(row=i, column=2, value=value)
    c2.font = body_font(); c2.border = thin_border()
    c2.fill = solid(fill); c2.alignment = left_align()

set_col_widths(ws, [32, 28])
print("  ✅  Sheet 1 — Overview")

# =============================================================================
# SHEET 2 — CONTINUOUS COLUMNS
# =============================================================================
ws = wb.create_sheet("Continuous Stats")
ws.sheet_view.showGridLines = False

HDR = ["Column", "Count", "Mean", "Std Dev", "Min",
       "P25", "Median", "P75", "Max", "Range",
       "Skewness", "Kurtosis", "CV %"]
write_banner(ws, 1, len(HDR), "Continuous Columns — Descriptive Statistics")
write_col_headers(ws, 2, HDR)

stats = df[continuous_cols].describe(percentiles=[0.25, 0.5, 0.75]).T
stats["range"]    = stats["max"] - stats["min"]
stats["skewness"] = df[continuous_cols].skew()
stats["kurtosis"] = df[continuous_cols].kurt()
stats["cv_pct"]   = (stats["std"] / stats["mean"].replace(0, np.nan) * 100)

for r, col in enumerate(continuous_cols, 3):
    f = row_fill(r)
    s = stats.loc[col]
    w_str(ws, r, 1,  col,                bold=True, fill_hex=f)
    w_num(ws, r, 2,  s["count"],         "#,##0",    f)
    w_num(ws, r, 3,  s["mean"],          "#,##0.00", f)
    w_num(ws, r, 4,  s["std"],           "#,##0.00", f)
    w_num(ws, r, 5,  s["min"],           "#,##0.00", f)
    w_num(ws, r, 6,  s["25%"],           "#,##0.00", f)
    w_num(ws, r, 7,  s["50%"],           "#,##0.00", f)
    w_num(ws, r, 8,  s["75%"],           "#,##0.00", f)
    w_num(ws, r, 9,  s["max"],           "#,##0.00", f)
    w_num(ws, r, 10, s["range"],         "#,##0.00", f)
    w_num(ws, r, 11, s["skewness"],      "0.000",    f)
    w_num(ws, r, 12, s["kurtosis"],      "0.000",    f)
    w_num(ws, r, 13, s["cv_pct"],        "0.00",     f)

set_col_widths(ws, [30,9,12,12,12,10,10,10,12,10,10,10,8])
ws.freeze_panes = "B3"
print("  ✅  Sheet 2 — Continuous Stats")

# =============================================================================
# SHEET 3 — BINARY FLAGS
# =============================================================================
ws = wb.create_sheet("Binary Flags")
ws.sheet_view.showGridLines = False

HDR = ["Column", "Total", "Count = 1", "Count = 0", "% True", "% False"]
write_banner(ws, 1, len(HDR), "Binary Flag Columns — Proportion Analysis")
write_col_headers(ws, 2, HDR)

flag_rows = []
for col in flag_cols:
    n1   = int(df[col].sum())
    n0   = int((df[col] == 0).sum())
    flag_rows.append((col, len(df), n1, n0, n1 / len(df), n0 / len(df)))
flag_rows.sort(key=lambda x: x[4], reverse=True)

for r, (col, total, n1, n0, pt, pf) in enumerate(flag_rows, 3):
    f = row_fill(r)
    w_str(ws, r, 1, col,   bold=True, fill_hex=f)
    w_num(ws, r, 2, total, "#,##0", f)
    w_num(ws, r, 3, n1,    "#,##0", f)
    w_num(ws, r, 4, n0,    "#,##0", f)
    c5 = w_num(ws, r, 5, pt, "0.0%")
    c5.fill = solid(GREEN_FILL if pt >= 0.50 else (RED_FILL if pt <= 0.10 else f))
    w_num(ws, r, 6, pf, "0.0%", f)

set_col_widths(ws, [28, 10, 12, 12, 10, 10])
ws.freeze_panes = "A3"
print("  ✅  Sheet 3 — Binary Flags")

# =============================================================================
# SHEET 4 — ONE-HOT DUMMIES (% True per category)
# =============================================================================
ws = wb.create_sheet("Dummy Variables")
ws.sheet_view.showGridLines = False

HDR = ["Original Variable", "Dummy Column", "Count = 1", "% of Rows"]
write_banner(ws, 1, len(HDR), "One-Hot Dummy Variables — Category Frequency")
write_col_headers(ws, 2, HDR)

def original_var(col):
    for p in DUMMY_PREFIXES:
        if col.startswith(p):
            return p.rstrip("_")
    return "other"

dummy_rows = [(original_var(c), c, int(df[c].sum()), df[c].mean())
              for c in dummy_cols]
dummy_rows.sort(key=lambda x: (x[0], -x[3]))

for r, (orig, col, cnt, pct) in enumerate(dummy_rows, 3):
    f = row_fill(r)
    w_str(ws, r, 1, orig, bold=True,  fill_hex=f)
    w_str(ws, r, 2, col,  bold=False, fill_hex=f)
    w_num(ws, r, 3, cnt,  "#,##0", f)
    w_num(ws, r, 4, pct,  "0.0%",  f)

set_col_widths(ws, [26, 36, 12, 10])
ws.freeze_panes = "A3"
print("  ✅  Sheet 4 — Dummy Variables")

# =============================================================================
# SHEET 5 — TARGET ANALYSIS
# =============================================================================
ws = wb.create_sheet("Target Analysis")
ws.sheet_view.showGridLines = False

write_banner(ws, 1, 4, f"Target Variable: {TARGET} — Feature-Level Analysis")

# Section A: target distribution
write_col_headers(ws, 2, ["Target Value", "Label", "Count", "% of Total"], fill=NAVY)
dist_rows = [
    (0, "Non-Holder", int((df[TARGET] == 0).sum()), (df[TARGET] == 0).mean()),
    (1, "Card Holder", int(df[TARGET].sum()),        df[TARGET].mean()),
]
for r, (val, label, cnt, pct) in enumerate(dist_rows, 3):
    f = row_fill(r)
    w_num(ws, r, 1, val,   "#,##0", f)
    w_str(ws, r, 2, label, fill_hex=f)
    w_num(ws, r, 3, cnt,   "#,##0", f)
    w_num(ws, r, 4, pct,   "0.0%",  f)

# Section B: mean of each continuous col by target
ws.cell(row=6, column=1, value="Mean by Target Class (Continuous Columns)").font = \
    Font(name="Arial", bold=True, size=10, color=NAVY)

write_col_headers(ws, 7, ["Column", "Mean (Target=0)", "Mean (Target=1)", "Difference", "% Diff"], fill=STEEL)

grp = df.groupby(TARGET)[continuous_cols].mean()
for r, col in enumerate(continuous_cols, 8):
    f  = row_fill(r)
    m0 = grp.loc[0, col] if 0 in grp.index else np.nan
    m1 = grp.loc[1, col] if 1 in grp.index else np.nan
    diff    = m1 - m0
    pct_diff = diff / m0 * 100 if m0 != 0 else np.nan
    w_str(ws, r, 1, col, bold=True, fill_hex=f)
    w_num(ws, r, 2, m0,      "#,##0.00", f)
    w_num(ws, r, 3, m1,      "#,##0.00", f)
    d_cell = w_num(ws, r, 4, diff,     "#,##0.00")
    d_cell.fill = solid(GREEN_FILL if diff > 0 else RED_FILL)
    w_num(ws, r, 5, pct_diff, "0.00",   f)

set_col_widths(ws, [30, 18, 18, 14, 10])
ws.freeze_panes = "A8"
print("  ✅  Sheet 5 — Target Analysis")

# =============================================================================
# SHEET 6 — TOP CORRELATIONS
# =============================================================================
ws = wb.create_sheet("Top Correlations")
ws.sheet_view.showGridLines = False

HDR = ["Rank", "Column A", "Column B", "Pearson r", "Abs r", "Strength"]
write_banner(ws, 1, len(HDR), "Top 30 Absolute Correlations — Continuous Features")
write_col_headers(ws, 2, HDR)

corr_matrix = df[continuous_cols].corr().abs()
upper = corr_matrix.where(np.triu(np.ones(corr_matrix.shape), k=1).astype(bool))
top30 = (upper.stack()
              .reset_index()
              .rename(columns={"level_0": "A", "level_1": "B", 0: "abs_r"})
              .sort_values("abs_r", ascending=False)
              .head(30))
top30["r"] = [df[row.A].corr(df[row.B]) for _, row in top30.iterrows()]
top30["strength"] = top30["abs_r"].apply(
    lambda x: "Very Strong" if x >= 0.80 else
              "Strong"      if x >= 0.60 else
              "Moderate"    if x >= 0.40 else "Weak"
)

STRENGTH_STYLE = {
    "Very Strong": ("1E8449", WHITE),
    "Strong":      ("58D68D", DARK_GRAY),
    "Moderate":    ("F9E79F", DARK_GRAY),
    "Weak":        ("FADBD8", DARK_GRAY),
}
for r, (_, row) in enumerate(top30.iterrows(), 3):
    f = row_fill(r)
    w_num(ws, r, 1, r - 2,        "#,##0",  f)
    w_str(ws, r, 2, row.A,        fill_hex=f)
    w_str(ws, r, 3, row.B,        fill_hex=f)
    w_num(ws, r, 4, row.r,        "0.0000", f)
    w_num(ws, r, 5, row.abs_r,    "0.0000", f)
    bg, fg = STRENGTH_STYLE[row.strength]
    c = ws.cell(row=r, column=6, value=row.strength)
    c.fill = solid(bg); c.font = body_font(bold=True, color=fg)
    c.alignment = center_align(); c.border = thin_border()

set_col_widths(ws, [7, 30, 30, 12, 12, 14])
ws.freeze_panes = "A3"
print("  ✅  Sheet 6 — Top Correlations")

# ── Save ──────────────────────────────────────────────────────────────────────
wb.save(OUTPUT_FILE)
print(f"\n🎉  Saved → {OUTPUT_FILE}")
